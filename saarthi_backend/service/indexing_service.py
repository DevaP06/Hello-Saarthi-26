"""Background material / transcript indexing service.

When a teacher uploads a material, this service:
1. Extracts text (PDF per page, TXT, DOCX, PPTX per slide)
2. Chunks it into ~500-character pieces (page/slide kept in metadata)
3. Embeds and upserts into the per-course FAISS index
4. Tags every chunk with material_id / material_title so retrieve() can filter

Indexing runs in a background thread so it never blocks the HTTP response.
Progress is tracked per material so the UI can poll
GET /api/courses/{course_id}/materials/{material_id}/index-status.

Status values: queued | running | done | error | unsupported | none
"""

import json
import re
import threading
import time
from collections import defaultdict
from pathlib import Path
from typing import Optional

from saarthi_backend.utils.logging import get_logger
from src.retrieval.schema import chunk_metadata

logger = get_logger(__name__)

# Anchor to the project root so indexes land in the same place retrieve()
# reads from, regardless of the directory the server was started in.
_PROJECT_ROOT = Path(__file__).resolve().parents[2]
_KB_ROOT = _PROJECT_ROOT / "knowledge_base" / "courses"
_VIDEO_KB_ROOT = _PROJECT_ROOT / "knowledge_base" / "videos"
_META_FILE = "index_meta.json"

_CHUNK_SIZE = 500       # characters per chunk
_CHUNK_OVERLAP = 80
_VIDEO_WINDOW_SEC = 45  # transcript window length for video chunks

INDEXABLE_EXTENSIONS = frozenset({".pdf", ".txt", ".md", ".docx", ".pptx"})


# ── Locks + status registry ──────────────────────────────────────────────────

_locks_guard = threading.Lock()
_locks: dict[str, threading.Lock] = defaultdict(threading.Lock)

_status_guard = threading.Lock()
_status: dict[int, dict] = {}


def _lock_for(key: str) -> threading.Lock:
    """One lock per index folder: concurrent uploads to the same course must not
    load-modify-save the same FAISS files at the same time (lost updates)."""
    with _locks_guard:
        return _locks[key]


def _set_status(material_id: int, status: str, message: str = "", **extra) -> None:
    with _status_guard:
        _status[material_id] = {
            "status": status,
            "message": message,
            "updatedAt": time.time(),
            **extra,
        }


def _read_meta(index_path: Path) -> dict:
    meta_file = index_path / _META_FILE
    if meta_file.exists():
        try:
            return json.loads(meta_file.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def _write_meta(index_path: Path, meta: dict) -> None:
    try:
        (index_path / _META_FILE).write_text(json.dumps(meta, indent=2), encoding="utf-8")
    except Exception as e:
        logger.warning("Could not write %s in %s: %s", _META_FILE, index_path, e)


def _backfill_meta(index_path: Path) -> dict:
    """Indexes built before index_meta.json existed: record which material_ids are
    already in the docstore so their status reads 'done' instead of 'none'."""
    meta = _read_meta(index_path)
    if meta.get("backfilled") or not (index_path / "index.pkl").exists():
        return meta
    with _lock_for(f"course:{index_path.parent.name}"):
        meta = _read_meta(index_path)
        if meta.get("backfilled"):
            return meta
        try:
            import pickle
            with open(index_path / "index.pkl", "rb") as f:
                docstore, _ = pickle.load(f)
            counts: dict[str, dict] = {}
            for doc in getattr(docstore, "_dict", {}).values():
                mid = (doc.metadata or {}).get("material_id")
                if mid is None:
                    continue
                entry = counts.setdefault(str(mid), {
                    "title": doc.metadata.get("material_title", ""),
                    "chunks": 0,
                    "indexed_at": None,
                })
                entry["chunks"] += 1
            materials = meta.setdefault("materials", {})
            for mid, entry in counts.items():
                materials.setdefault(mid, entry)
        except Exception as e:
            logger.warning("Could not backfill %s for %s: %s", _META_FILE, index_path, e)
        meta["backfilled"] = True
        _write_meta(index_path, meta)
        return meta


def get_material_index_status(course_id: int, material_id: int) -> dict:
    """Current indexing status for a material. Survives restarts via index_meta.json."""
    with _status_guard:
        live = _status.get(material_id)
    if live:
        return {k: v for k, v in live.items() if k != "updatedAt"}
    meta = _backfill_meta(_KB_ROOT / str(course_id) / "vector_store")
    entry = (meta.get("materials") or {}).get(str(material_id))
    if entry:
        return {"status": "done", "message": "Indexed", "chunks": entry.get("chunks", 0)}
    return {"status": "none", "message": "Not indexed"}


def _invalidate_cached_answers(course_id: int | None = None, video_id: int | None = None) -> None:
    """Course/video content changed: cached answers for that scope may be stale."""
    try:
        from src.cache.semantic_cache import get_cache
        if course_id is not None:
            n = get_cache().invalidate_course(course_id)
        else:
            n = get_cache().invalidate_video(video_id)
        if n:
            logger.info("Semantic cache: dropped %d entries (course=%s video=%s)", n, course_id, video_id)
    except Exception as e:
        logger.debug("Semantic cache invalidation skipped: %s", e)


def is_indexable(file_path: Path | str) -> bool:
    return Path(file_path).suffix.lower() in INDEXABLE_EXTENSIONS


# ── Text extraction ──────────────────────────────────────────────────────────

def _extract_pdf_pages(path: Path) -> list[tuple[Optional[int], str]]:
    """(page_number, text) per page using pypdf (no OCR needed for text PDFs)."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        out = []
        for i, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                out.append((i, text))
        return out
    except Exception as e:
        logger.warning("PDF text extraction failed for %s: %s", path, e)
        return []


def _extract_text_from_pdf(path: Path) -> str:
    """Whole-document PDF text (kept for callers that don't need page numbers)."""
    return "\n\n".join(text for _, text in _extract_pdf_pages(path))


def _extract_text_from_txt(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        logger.warning("Text extraction failed for %s: %s", path, e)
        return ""


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        logger.warning("DOCX extraction failed for %s: %s", path, e)
        return ""


def _extract_pptx_slides(path: Path) -> list[tuple[Optional[int], str]]:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text.strip())
            if texts:
                out.append((i, "\n".join(texts)))
        return out
    except Exception as e:
        logger.warning("PPTX extraction failed for %s: %s", path, e)
        return []


def _extract_segments(path: Path) -> list[tuple[Optional[int], str]]:
    """Return (page_or_slide, text) segments. Page is None for flowing text."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf_pages(path)
    if ext in (".txt", ".md"):
        return [(None, _extract_text_from_txt(path))]
    if ext == ".docx":
        return [(None, _extract_docx(path))]
    if ext == ".pptx":
        return _extract_pptx_slides(path)
    return []


def _chunk_text(text: str, chunk_size: int = _CHUNK_SIZE, overlap: int = _CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks."""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


# ── FAISS upsert ─────────────────────────────────────────────────────────────

def _upsert_index(
    index_path: Path,
    texts: list[str],
    metadatas: list[dict],
    drop_where: Optional[tuple[str, object]] = None,
) -> None:
    """Add texts to the FAISS index at index_path (creating it if needed).

    drop_where=(key, value) first removes existing chunks whose metadata[key]==value,
    so re-indexing a material replaces its chunks instead of duplicating them.
    If the stored index was built with a different embedding size (EMBEDDING_MODEL
    changed), the existing chunks are re-embedded with the current model instead of
    failing with a dimension mismatch.
    """
    from langchain_community.vectorstores import FAISS
    from src.utils.llm import embeddings as make_embeddings, embedding_model_name

    emb = make_embeddings()
    vectors = emb.embed_documents(texts)
    dim = len(vectors[0])
    index_path.mkdir(parents=True, exist_ok=True)

    vs = None
    if (index_path / "index.faiss").exists():
        vs = FAISS.load_local(str(index_path), emb, allow_dangerous_deserialization=True)
        if drop_where is not None:
            key, value = drop_where
            stale = [i for i, d in vs.docstore._dict.items() if d.metadata.get(key) == value]
            if stale:
                vs.delete(stale)
        if vs.index.d != dim:
            old_docs = list(vs.docstore._dict.values())
            logger.warning(
                "Embedding size changed (%d -> %d) for %s; re-embedding %d existing chunks",
                vs.index.d, dim, index_path, len(old_docs),
            )
            if old_docs:
                old_vectors = emb.embed_documents([d.page_content for d in old_docs])
                texts = [d.page_content for d in old_docs] + texts
                vectors = old_vectors + vectors
                metadatas = [d.metadata for d in old_docs] + metadatas
            vs = None

    if vs is None:
        vs = FAISS.from_embeddings(list(zip(texts, vectors)), emb, metadatas=metadatas)
    else:
        vs.add_embeddings(list(zip(texts, vectors)), metadatas=metadatas)
    vs.save_local(str(index_path))

    meta = _read_meta(index_path)
    meta["embedding_model"] = embedding_model_name()
    meta["dim"] = dim
    _write_meta(index_path, meta)


def _index_document_sync(
    file_path: Path | str,
    material_title: str,
    course_id: int,
    material_id: int,
) -> bool:
    """Blocking: extract, chunk, embed, and upsert into course FAISS index."""
    file_path = Path(file_path)
    if not is_indexable(file_path):
        _set_status(material_id, "unsupported", f"{file_path.suffix or 'This file type'} can't be indexed")
        logger.info("Skipping indexing for unsupported type %s (material: %s)", file_path.suffix, material_title)
        return False

    lock = _lock_for(f"course:{course_id}")
    with lock:
        _set_status(material_id, "running", "Reading the document…")
        try:
            segments = _extract_segments(file_path)
            source_tag = material_title.lower().replace(" ", "_")
            texts: list[str] = []
            metadatas: list[dict] = []
            for page, seg_text in segments:
                for chunk in _chunk_text(seg_text or ""):
                    metadatas.append(chunk_metadata(
                        "course_material", chunk,
                        source=source_tag,
                        material_title=material_title,
                        course_id=course_id,
                        material_id=material_id,
                        page=page,
                        chunk_index=len(texts),
                    ))
                    texts.append(chunk)

            if sum(len(t) for t in texts) < 50:
                _set_status(material_id, "error", "No readable text found (scanned PDF?)")
                logger.info("No extractable text in %s, skipping indexing", file_path.name)
                return False

            _set_status(material_id, "running", f"Embedding {len(texts)} chunks…")
            index_path = _KB_ROOT / str(course_id) / "vector_store"
            _upsert_index(index_path, texts, metadatas, drop_where=("material_id", material_id))

            meta = _read_meta(index_path)
            meta.setdefault("materials", {})[str(material_id)] = {
                "title": material_title,
                "chunks": len(texts),
                "indexed_at": time.time(),
            }
            _write_meta(index_path, meta)
            _set_status(material_id, "done", "Indexed", chunks=len(texts))
            _invalidate_cached_answers(course_id=course_id)
            logger.info(
                "Indexed %d chunks for material '%s' (course %d) into %s",
                len(texts), material_title, course_id, index_path,
            )
            return True
        except Exception as e:
            msg = str(e)
            if "Connection" in type(e).__name__ or "connect" in msg.lower():
                msg = "Embedding service unreachable (is Ollama / the API running?)"
            _set_status(material_id, "error", f"Indexing failed: {msg[:200]}")
            logger.error("FAISS indexing failed for material '%s': %s", material_title, e)
            return False


def index_material_background(
    file_path: Path | str,
    material_title: str,
    course_id: int,
    material_id: int,
) -> None:
    """Fire-and-forget: index a material in a background thread.

    The status is set to 'queued' before returning so a poll that arrives
    right after the HTTP response never sees 'none'.
    """
    file_path = Path(file_path)
    if not is_indexable(file_path):
        _set_status(material_id, "unsupported", f"{file_path.suffix or 'This file type'} can't be indexed")
        return
    _set_status(material_id, "queued", "Waiting to index…")

    def _run():
        logger.info("Background indexing started for '%s' (course %d)", material_title, course_id)
        ok = _index_document_sync(file_path, material_title, course_id, material_id)
        if ok:
            logger.info("Background indexing complete for '%s'", material_title)
        else:
            logger.warning("Background indexing skipped/failed for '%s'", material_title)

    t = threading.Thread(target=_run, daemon=True)
    t.start()


# ── Video transcripts ────────────────────────────────────────────────────────

def _extract_youtube_id(url: str) -> str | None:
    """Pull the 11-char video ID from any YouTube URL or embed URL."""
    m = re.search(r"(?:v=|/v/|youtu\.be/|/embed/|/shorts/)([A-Za-z0-9_-]{11})", url or "")
    return m.group(1) if m else None


def fetch_youtube_transcript_segments(url: str) -> list[dict]:
    """Each item: {text, start_sec, duration}."""
    video_id = _extract_youtube_id(url)
    if not video_id:
        return []
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
        segs = YouTubeTranscriptApi().fetch(video_id)
        out = []
        for s in segs:
            text = getattr(s, "text", "") or ""
            if not text:
                continue
            out.append({
                "text": text,
                "start_sec": int(getattr(s, "start", 0) or 0),
                "duration": int(getattr(s, "duration", 0) or 0),
            })
        return out
    except Exception as e:
        logger.warning("YouTube transcript fetch failed for %s: %s", video_id, e)
        return []


def fetch_youtube_transcript(url: str) -> str:
    """Plain text transcript (legacy). Prefer segments for timestamps."""
    return " ".join(s["text"] for s in fetch_youtube_transcript_segments(url))


def _index_video_segments_sync(video_id: int, video_title: str, segs: list[dict]) -> bool:
    """Blocking: group transcript segments into ~45s windows (start_sec/end_sec kept
    in metadata for citations) and store them in the per-video FAISS index."""
    texts: list[str] = []
    metadatas: list[dict] = []
    source_tag = video_title.lower().replace(" ", "_")
    buf: list[str] = []
    win_start: int | None = None
    win_end = 0

    def _flush():
        if buf and win_start is not None:
            text = " ".join(buf)
            metadatas.append(chunk_metadata(
                "video", text,
                source=source_tag,
                video_title=video_title,
                video_id=video_id,
                start_sec=win_start,
                end_sec=win_end,
                chunk_index=len(texts),
            ))
            texts.append(text)

    for s in segs:
        start = int(s.get("start_sec") or 0)
        if win_start is None:
            win_start = start
        buf.append(s.get("text", ""))
        win_end = start + int(s.get("duration") or 0)
        if win_end - win_start >= _VIDEO_WINDOW_SEC:
            _flush()
            buf, win_start = [], None
    _flush()

    if sum(len(t) for t in texts) < 50:
        logger.info("Transcript too short to index for video %d", video_id)
        return False

    with _lock_for(f"video:{video_id}"):
        try:
            index_path = _VIDEO_KB_ROOT / str(video_id) / "vector_store"
            # A video has exactly one transcript: replace, don't append.
            if index_path.exists():
                import shutil
                shutil.rmtree(index_path, ignore_errors=True)
            _upsert_index(index_path, texts, metadatas)
            logger.info("Indexed %d transcript windows for video %d ('%s')", len(texts), video_id, video_title)
            _invalidate_cached_answers(video_id=video_id)
            return True
        except Exception as e:
            logger.error("FAISS transcript indexing failed for video %d: %s", video_id, e)
            return False


def auto_index_youtube_video(video_id: int, video_title: str, url: str, embed_url: str | None = None) -> None:
    """
    Fire-and-forget: if the video is a YouTube video and has no index yet,
    fetch its transcript and index it automatically.
    """
    index_path = _VIDEO_KB_ROOT / str(video_id) / "vector_store"
    if (index_path / "index.faiss").exists():
        return  # already indexed

    yt_url = embed_url or url
    if not _extract_youtube_id(yt_url) and not _extract_youtube_id(url):
        return  # not a YouTube video

    def _run():
        segs = fetch_youtube_transcript_segments(yt_url) or fetch_youtube_transcript_segments(url)
        if segs:
            logger.info("Auto-fetched YouTube transcript for video %d (%d segs)", video_id, len(segs))
            _index_video_segments_sync(video_id, video_title, segs)
        else:
            logger.info("No YouTube transcript available for video %d", video_id)

    t = threading.Thread(target=_run, daemon=True)
    t.start()


def _index_video_transcript_sync(video_id: int, video_title: str, transcript_text: str) -> bool:
    """Blocking: chunk, embed, and store a pasted transcript into the per-video FAISS index."""
    if not transcript_text or len(transcript_text.strip()) < 50:
        logger.info("Transcript too short to index for video %d", video_id)
        return False

    chunks = _chunk_text(transcript_text)
    source_tag = video_title.lower().replace(" ", "_")
    metadatas = [
        chunk_metadata("video", c, source=source_tag, video_title=video_title, video_id=video_id, chunk_index=i)
        for i, c in enumerate(chunks)
    ]
    with _lock_for(f"video:{video_id}"):
        try:
            index_path = _VIDEO_KB_ROOT / str(video_id) / "vector_store"
            if index_path.exists():
                import shutil
                shutil.rmtree(index_path, ignore_errors=True)
            _upsert_index(index_path, chunks, metadatas)
            logger.info("Indexed %d transcript chunks for video %d ('%s')", len(chunks), video_id, video_title)
            _invalidate_cached_answers(video_id=video_id)
            return True
        except Exception as e:
            logger.error("FAISS transcript indexing failed for video %d: %s", video_id, e)
            return False


def index_video_transcript_background(video_id: int, video_title: str, transcript_text: str) -> None:
    """Fire-and-forget: index a video transcript in a background thread."""
    def _run():
        logger.info("Background transcript indexing started for video %d", video_id)
        ok = _index_video_transcript_sync(video_id, video_title, transcript_text)
        if ok:
            logger.info("Transcript indexing complete for video %d", video_id)
        else:
            logger.warning("Transcript indexing skipped/failed for video %d", video_id)

    t = threading.Thread(target=_run, daemon=True)
    t.start()


def fetch_video_transcript_context(query: str, video_id: int, video_title: str, k: int = 5) -> str:
    """Search the per-video FAISS index for chunks relevant to the query."""
    try:
        from langchain_community.vectorstores import FAISS
        from src.utils.llm import embeddings as make_embeddings

        index_path = _VIDEO_KB_ROOT / str(video_id) / "vector_store"
        if not (index_path / "index.faiss").exists():
            return ""

        vs = FAISS.load_local(str(index_path), make_embeddings(), allow_dangerous_deserialization=True)
        docs = vs.similarity_search(f"{video_title} {query}", k=k)
        chunks = [d.page_content.strip() for d in docs if d.page_content.strip()]
        return "\n\n---\n\n".join(chunks[:k])
    except Exception as e:
        logger.warning("Video transcript context fetch failed for video %d: %s", video_id, e)
        return ""


def delete_video_transcript_index(video_id: int) -> None:
    """Remove the per-video FAISS index when a video is deleted."""
    import shutil
    index_path = _VIDEO_KB_ROOT / str(video_id)
    if index_path.exists():
        shutil.rmtree(str(index_path), ignore_errors=True)
        logger.info("Deleted transcript index for video %d", video_id)


def delete_material_from_index(
    material_title: str,
    course_id: int,
    material_id: int,
) -> None:
    """Remove a material's chunks from the course FAISS index on deletion."""
    with _status_guard:
        _status.pop(material_id, None)

    def _run():
        with _lock_for(f"course:{course_id}"):
            try:
                from langchain_community.vectorstores import FAISS
                from src.utils.llm import embeddings as make_embeddings

                index_path = _KB_ROOT / str(course_id) / "vector_store"
                if not (index_path / "index.faiss").exists():
                    return

                vs = FAISS.load_local(
                    str(index_path), make_embeddings(), allow_dangerous_deserialization=True
                )
                ids_to_remove = [
                    doc_id for doc_id, doc in vs.docstore._dict.items()
                    if doc.metadata.get("material_id") == material_id
                ]
                if ids_to_remove:
                    vs.delete(ids_to_remove)
                    vs.save_local(str(index_path))
                    logger.info(
                        "Removed %d chunks for material_id=%d from course %d index",
                        len(ids_to_remove), material_id, course_id,
                    )
                meta = _read_meta(index_path)
                if (meta.get("materials") or {}).pop(str(material_id), None) is not None:
                    _write_meta(index_path, meta)
                _invalidate_cached_answers(course_id=course_id)
            except Exception as e:
                logger.warning("Failed to remove material from index: %s", e)

    t = threading.Thread(target=_run, daemon=True)
    t.start()
